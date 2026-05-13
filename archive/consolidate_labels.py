import openpyxl

EXCEL = r'C:\MyApps\ShoppingCart\Combined Justification 8.05.2026.xlsx'

CORRECTIONS = {
    # Bakery - sentence-case fix for all 19 items
    'B1':'Apple danish','B2':'Bagel plain','B3':'BBQ cheese & bacon roll',
    'B4':'Butter croissant','B5':'Cheese & bacon roll','B6':'Choc Chip hot cross bun',
    'B7':'Cinnamon scroll','B8':'Double choc muffin','B9':'Fruit filled tart',
    'B10':'Ham & cheese croissant','B11':'High fibre white bread','B12':'Jam filled doughnut',
    'B13':'Multigrain bread','B14':'Rye bread','B15':'Wholemeal wrap',
    'B16':'Wholemeal bread','B17':'Wholemeal english muffin','B18':'Wholemeal pizza base',
    'B19':'Wholemeal roll',
    # Dairy - all 28 confirmed against qualtrics tab
    'DEF1':'Reduced fat custard','DEF2':'Reduced fat milk','DEF3':'Ricotta cheese',
    'DEF4':'Soy milk','DEF5':'Protein greek yoghurt','DEF6':'Eggs',
    'DEF7':'Cottage cheese','DEF8':'Reduced fat cheddar cheese','DEF9':'Fresh napoli sauce',
    'DEF10':'Tofu','DEF11':'Leg ham','DEF12':'Coconut yoghurt','DEF13':'Salami',
    'DEF14':'Roast Pastrami Sliced','DEF15':'Chicken breast sliced',
    'DEF16':'Fresh carbonara pasta sauce','DEF17':'Middle Bacon','DEF18':'Silverside sliced',
    'DEF19':'Chocolate chip yoghurt','DEF20':'Maple Bacon','DEF21':'Chicken shredded',
    'DEF22':'Kabana','DEF23':'Prosciutto Sliced','DEF24':'Turkey breast sliced',
    'DEF25':'High protein chocolate pudding','DEF26':'Sliced Pancetta',
    'DEF27':'Smiley Fritz','DEF28':'Short cut bacon',
    # Drinks - stale Final labels corrected
    'DRK1':'Natural spring water','DRK2':'Lightly sparkling water',
    'DRK4':'Orange juice','DRK7':'Electrolyte sport drink','DRK11':'Protein water berry',
    # Snacks
    'SNA17':'Chicken flavoured potato chips',
    # Fruit - all 20 populated from qualtrics (13 were empty, 7 moved from Final label)
    'FRU1':'Nectarine','FRU2':'Grape','FRU3':'Watermelon',
    'FRU4':'Strawberry','FRU5':'Blueberry','FRU6':'Pear','FRU7':'Pineapple',
    'FRU8':'Apple','FRU9':'Banana','FRU10':'Orange','FRU11':'Kiwi fruit',
    'FRU12':'Canned peache','FRU13':'Canned fruit salad','FRU14':'Dried apple',
    'FRU15':'Dried mango','FRU16':'Banana Chip',
    'FRU17':'Freeze dried fruit dipped in chocolate',
    'FRU18':'Fruit Jelly Cup','FRU19':'Fruit string','FRU20':'Fruit roll up',
    # Meat Seafood - stale Final labels corrected
    'MEA1':'Lean beef mince','MEA2':'Steak','MEA5':'Lean pork mince','MEA9':'Basa fillet',
    # Freezer - 2 missing populated, 1 stale corrected
    'FRZ12':'Chocolate icecream','FRZ21':'Strawberry icecream','FRZ22':'Vanilla icecream',
    # Pantry - stale Final labels corrected
    'PAN8':'Rice','PAN9':'Spaghetti','PAN16':'Black beans reduced salt',
    'PAN18':'Red lentils dried','PAN21':'Tuna in springwater','PAN22':'Salmon in springwater',
    'PAN26':'Beetroot canned',
    # Ready to eat
    'RTE2':'Green salad','RTE18':'Meat pie',
    # Vegetables
    'VEG13':'Bok Choy',
}

TABS_WITH_FINAL = ['Drinks','Snacks','Fruit ','Meat Seafood','Freezer','Pantry','Ready to eat ','Vegetables']
TABS_LABEL_ONLY = ['Bakery  ','Dairy Eggs Fridge']

wb = openpyxl.load_workbook(EXCEL)

total_moved = 0
total_corrected = 0
total_cols_deleted = 0

for tab in TABS_WITH_FINAL + TABS_LABEL_ONLY:
    ws = wb[tab]
    header = [str(c.value).strip() if c.value else '' for c in ws[1]]
    h_low  = [h.lower() for h in header]

    fl_col_idx = next((i+1 for i,v in enumerate(h_low) if 'final' in v), None)

    moved = 0
    corrected = 0

    for row in ws.iter_rows(min_row=2):
        iid_cell = row[0]
        if not iid_cell.value:
            continue
        iid    = str(iid_cell.value).strip()
        d_cell = row[3]
        fl_cell = row[fl_col_idx - 1] if fl_col_idx else None

        fl_val = str(fl_cell.value).strip() if fl_cell and fl_cell.value else None

        if iid in CORRECTIONS:
            d_cell.value = CORRECTIONS[iid]
            corrected += 1
        elif fl_val:
            d_cell.value = fl_val
            moved += 1

    total_moved += moved
    total_corrected += corrected

    if fl_col_idx:
        ws.delete_cols(fl_col_idx)
        total_cols_deleted += 1
        print(f'  {tab.strip():20} | moved {moved:2d} | corrected {corrected:2d} | deleted col {fl_col_idx}')
    else:
        print(f'  {tab.strip():20} |          | corrected {corrected:2d} | (no Final label col)')

wb.save(EXCEL)
print()
print(f'Done. {total_corrected} corrections, {total_moved} Final labels moved, {total_cols_deleted} cols deleted.')
